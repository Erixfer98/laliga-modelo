"""
backtest.py — Backtest walk-forward del Poisson + Dixon-Coles (modelo.py) sobre TODAS las metricas de la BBDD.

Como funciona (igual que si Kuota se hubiera consultado la manana de cada jornada):
  1. Para cada fecha con partidos, el modelo se ajusta SOLO con los partidos anteriores a esa fecha
     (misma recencia XI y mismo encogimiento K_ENCOGE de modelo.py). Nada del futuro entra.
  2. Se calculan las probabilidades de cada mercado y se comparan con lo que paso.
  3. Calificacion = la misma de la app: 0.6 x prob. Poisson + 0.4 x cumplimiento en los ultimos 5 partidos
     de cada equipo (N_HIST).

Mercados por metrica (los mismos que muestra la app):
  goles            Total O/U 0.5-4.5 · Ambos anotan · Local y Visitante O/U 0.5-2.5
  goles_1t/2t      Total O/U 0.5-2.5 · Local y Visitante O/U 0.5-1.5
  rojas            Total O/U 0.5-1.5 · Local y Visitante O/U 0.5
  tiros, tiros_a_puerta, corners, faltas, amarillas
                   Total O/U en la linea central del partido (c = round(lambda) - 0.5) y c-2 … c+2
                   Local y Visitante O/U en su linea central y ±1 · Mayor numero (quien hace mas)

Uso: python backtest.py                     -> lee datos/bbdd_<liga>.csv de las 6 ligas
     python backtest.py a.xlsx b.xlsx ...    -> usa esos archivos (hoja "partidos"); la liga sale del nombre
Salida: backtest_kuota.xlsx (resumenes) + backtest_predicciones.csv (una fila por partido, metrica, mercado y lado)
"""

import os
import re
import sys
from multiprocessing import Pool

import numpy as np
import pandas as pd

import modelo as mo

LIGAS = ["laliga", "premier", "seriea", "bundesliga", "ligue1", "ligamx"]
NOMBRE_LIGA = {"laliga": "La Liga", "premier": "Premier League", "seriea": "Serie A",
               "bundesliga": "Bundesliga", "ligue1": "Ligue 1", "ligamx": "Liga MX"}
METRICAS = list(mo.METRICAS)    # las 9 de modelo.py
GOL = ("goles", "goles_1t", "goles_2t")
MAYOR_NUMERO = ("tiros", "tiros_a_puerta", "corners", "faltas", "amarillas")
N_HIST = 5                      # ultimos N partidos de cada equipo para la parte historica de la calificacion
PJ_MIN = mo.PJ_MIN              # partidos en su condicion (local en casa / visitante fuera) para "datos suficientes"
NIVELES = ((0.78, "Excelente"), (0.68, "Buena"), (0.56, "Regular"), (0.45, "Mala"), (0.0, "Pésima"))
TRAMOS = [(0.5, 0.6), (0.6, 0.7), (0.7, 0.8), (0.8, 0.9), (0.9, 1.01)]
ORDEN_CAL = [t for _, t in NIVELES]


# ------------------------------------------------------------------ datos
def cargar(ruta: str) -> pd.DataFrame:
    df = pd.read_excel(ruta, sheet_name="partidos") if ruta.lower().endswith(".xlsx") else pd.read_csv(ruta)
    df["fecha"] = pd.to_datetime(df["fecha_dt"], dayfirst=True)
    return df.sort_values(["fecha", "equipo_local_txt"]).reset_index(drop=True)


def liga_de(ruta: str) -> str:
    base = os.path.basename(ruta).lower()
    for lg in LIGAS:
        if lg in base:
            return lg
    return re.sub(r"\.(xlsx|csv)$", "", base)


def lambdas_seguras(f: pd.DataFrame, local: str, visitante: str) -> tuple:
    """Como mo.lambdas, pero un equipo sin historial (recien ascendido en su 1er partido) vale promedio de liga."""
    def g(eq, col):
        return float(f.loc[eq, col]) if eq in f.index else 1.0
    lam_l = g(local, "ataque_casa") * g(visitante, "defensa_fuera") * f.attrs["prom_casa"]
    lam_v = g(visitante, "ataque_fuera") * g(local, "defensa_casa") * f.attrs["prom_fuera"]
    return lam_l, lam_v


def calificar(p: float, tasa: float) -> str:
    s = 0.6 * p + 0.4 * tasa
    for lim, txt in NIVELES:
        if s >= lim:
            return txt
    return "Pésima"


def centro(lam: float, met: str) -> float:
    """Linea central por defecto de la app (centro_defecto de app.py)."""
    if met == "goles":
        return 2.5
    if met in GOL:
        return 1.5
    if met == "rojas":
        return 0.5
    return max(round(lam) - 0.5, 0.5)


def lineas(met: str, lam_total: float, lam_l: float, lam_v: float) -> tuple:
    """(lineas total, lineas local, lineas visitante) y etiqueta relativa a la linea central."""
    if met == "goles":
        return [(x, str(x)) for x in (0.5, 1.5, 2.5, 3.5, 4.5)], [(x, str(x)) for x in (0.5, 1.5, 2.5)], [(x, str(x)) for x in (0.5, 1.5, 2.5)]
    if met in GOL:
        return [(x, str(x)) for x in (0.5, 1.5, 2.5)], [(x, str(x)) for x in (0.5, 1.5)], [(x, str(x)) for x in (0.5, 1.5)]
    if met == "rojas":
        return [(x, str(x)) for x in (0.5, 1.5)], [(0.5, "0.5")], [(0.5, "0.5")]
    ct, cl, cv = centro(lam_total, met), centro(lam_l, met), centro(lam_v, met)
    tot = [(ct + d, f"c{d:+d}" if d else "c") for d in (-2, -1, 0, 1, 2) if ct + d > 0]
    loc = [(cl + d, f"c{d:+d}" if d else "c") for d in (-1, 0, 1) if cl + d > 0]
    vis = [(cv + d, f"c{d:+d}" if d else "c") for d in (-1, 0, 1) if cv + d > 0]
    return tot, loc, vis


# ------------------------------------------------------------------ predicciones partido a partido
def predecir_liga(args) -> pd.DataFrame:
    ruta, metricas = args
    liga = liga_de(ruta)
    df = cargar(ruta)
    filas = []
    for fecha in sorted(df["fecha"].unique()):
        hist = df[df["fecha"] < fecha].copy()
        if hist.empty:
            continue
        hist["peso"] = np.exp(-mo.XI * (fecha - hist["fecha"]).dt.days)
        fuerzas = {met: mo.fuerzas(hist, met) for met in metricas}
        n_casa, n_fuera = hist["equipo_local_txt"].value_counts(), hist["equipo_visitante_txt"].value_counts()

        for _, p in df[df["fecha"] == fecha].iterrows():
            local, vis = p["equipo_local_txt"], p["equipo_visitante_txt"]
            base = {"liga": NOMBRE_LIGA.get(liga, liga), "temporada": p["temporada_txt"], "jornada": p["jornada_val"],
                    "fecha": fecha, "local": local, "visitante": vis,
                    "pj_casa_local": int(n_casa.get(local, 0)), "pj_fuera_visitante": int(n_fuera.get(vis, 0))}
            base["datos_suficientes"] = base["pj_casa_local"] >= PJ_MIN and base["pj_fuera_visitante"] >= PJ_MIN

            for met in metricas:
                col_l, col_v, dc = mo.METRICAS[met]
                xl, xv = int(p[col_l]), int(p[col_v])
                lam_l, lam_v = lambdas_seguras(fuerzas[met], local, vis)
                m = mo.matriz(lam_l, lam_v, mo.K_MAX[met], dc)
                k = m.shape[0]
                tot = np.add.outer(np.arange(k), np.arange(k))
                hl, hv = mo.ultimos_n(hist, local, met, N_HIST), mo.ultimos_n(hist, vis, met, N_HIST)
                nh = len(hl) + len(hv)
                b = {**base, "metrica": met, "real_local": xl, "real_visitante": xv, "real_total": xl + xv,
                     "lambda_local": lam_l, "lambda_visitante": lam_v, "lambda_total": lam_l + lam_v}

                def add(grupo, mercado, linea, lado, prob, cumple_l, cumple_v, acierto):
                    tasa = (np.sum(cumple_l) + np.sum(cumple_v)) / nh if nh else 0.0
                    prob = float(min(max(prob, 0.0), 1.0))
                    filas.append({**b, "grupo": grupo, "mercado": mercado, "linea": linea, "lado": lado,
                                  "prob_poisson": prob, "tasa_ult5": float(tasa), "calificacion": calificar(prob, tasa),
                                  "acierto": int(acierto)})

                lt, ll, lv = lineas(met, lam_l + lam_v, lam_l, lam_v)
                for ln, et in lt:
                    po = m[tot > ln].sum()
                    add("Total", f"Total Over {et}", ln, "Over", po, hl.total > ln, hv.total > ln, xl + xv > ln)
                    add("Total", f"Total Under {et}", ln, "Under", 1 - po, hl.total < ln, hv.total < ln, xl + xv < ln)
                if met == "goles":
                    pb = m[1:, 1:].sum()
                    bl, bv = (hl.a_favor > 0) & (hl.en_contra > 0), (hv.a_favor > 0) & (hv.en_contra > 0)
                    add("Ambos anotan", "Ambos anotan: Sí", 0.5, "Over", pb, bl, bv, xl > 0 and xv > 0)
                    add("Ambos anotan", "Ambos anotan: No", 0.5, "Under", 1 - pb, ~bl, ~bv, xl == 0 or xv == 0)
                for ln, et in ll:
                    po = mo.prob_over(lam_l, ln)
                    add("Local", f"Local Over {et}", ln, "Over", po, hl.a_favor > ln, hv.en_contra > ln, xl > ln)
                    add("Local", f"Local Under {et}", ln, "Under", 1 - po, hl.a_favor < ln, hv.en_contra < ln, xl < ln)
                for ln, et in lv:
                    po = mo.prob_over(lam_v, ln)
                    add("Visitante", f"Visitante Over {et}", ln, "Over", po, hl.en_contra > ln, hv.a_favor > ln, xv > ln)
                    add("Visitante", f"Visitante Under {et}", ln, "Under", 1 - po, hl.en_contra < ln, hv.a_favor < ln, xv < ln)
                if met in MAYOR_NUMERO:
                    p1, px, p2 = np.tril(m, -1).sum(), np.trace(m), np.triu(m, 1).sum()
                    gl, gv = hl.a_favor > hl.en_contra, hv.a_favor > hv.en_contra
                    add("Mayor número", "Más: local", 0.5, "Over", p1, gl, hv.a_favor < hv.en_contra, xl > xv)
                    add("Mayor número", "Más: empate", 0.5, "Over", px, hl.a_favor == hl.en_contra, hv.a_favor == hv.en_contra, xl == xv)
                    add("Mayor número", "Más: visitante", 0.5, "Over", p2, hl.a_favor < hl.en_contra, gv, xl < xv)
    return pd.DataFrame(filas)


# ------------------------------------------------------------------ resumenes
def picks(pred: pd.DataFrame) -> pd.DataFrame:
    """Una fila por partido y linea: el lado que el modelo prefiere (prob >= 50%). 'Mayor numero' no entra
    (tres opciones, no dos)."""
    over = pred[(pred["lado"] == "Over") & (pred["grupo"] != "Mayor número")].copy()
    over["pick"] = np.where(over["prob_poisson"] >= 0.5, "Over", "Under")
    over["prob_pick"] = np.where(over["pick"] == "Over", over["prob_poisson"], 1 - over["prob_poisson"])
    over["acierto_pick"] = np.where(over["pick"] == "Over", over["acierto"], 1 - over["acierto"])
    over["ocurrio_over"] = over["acierto"]
    over["mercado"] = over["mercado"].str.replace(" Over ", " ", regex=False).str.replace(": Sí", "", regex=False)
    return over


def ingenuo(g: pd.DataFrame) -> float:
    """Acierto de apostar siempre al lado mas frecuente de cada linea en cada liga (sin modelo)."""
    if len(g) == 0:
        return np.nan
    hits = sum(max(s.sum(), len(s) - s.sum()) for _, s in g.groupby(["liga", "metrica", "mercado"])["ocurrio_over"])
    return hits / len(g)


def resumen_mercados(pk: pd.DataFrame) -> pd.DataFrame:
    filas = []
    for (met, grupo, mercado), g in pk.groupby(["metrica", "grupo", "mercado"], sort=False):
        base = g["ocurrio_over"].mean()
        filas.append({
            "metrica": met, "grupo": grupo, "mercado": mercado, "partidos": len(g),
            "over_real_pct": base, "pct_picks_over": (g["pick"] == "Over").mean(),
            "acierto_modelo": g["acierto_pick"].mean(), "acierto_sin_modelo": ingenuo(g),
            "prob_media_pick": g["prob_pick"].mean(),
            "brier_modelo": ((g["prob_poisson"] - g["ocurrio_over"]) ** 2).mean(),
            "brier_sin_modelo": ((base - g["ocurrio_over"]) ** 2).mean(),
        })
    r = pd.DataFrame(filas)
    r["mejora_vs_sin_modelo"] = r["acierto_modelo"] - r["acierto_sin_modelo"]
    r["habilidad_brier"] = 1 - r["brier_modelo"] / r["brier_sin_modelo"]
    r["metrica"] = pd.Categorical(r["metrica"], METRICAS, ordered=True)
    return r.sort_values(["metrica", "grupo"]).reset_index(drop=True)


def calibracion(pk: pd.DataFrame, por=None) -> pd.DataFrame:
    filas = []
    grupos = [((), pk)] if por is None else list(pk.groupby(por, sort=False))
    for clave, g in grupos:
        clave = clave if isinstance(clave, tuple) else (clave,)
        etiquetas = dict(zip(por if isinstance(por, list) else ([por] if por else []), clave))
        for lo, hi in TRAMOS:
            s = g[(g["prob_pick"] >= lo) & (g["prob_pick"] < hi)]
            if len(s) == 0:
                continue
            filas.append({**etiquetas, "tramo_prob": f"{lo:.0%}–{min(hi, 1):.0%}", "picks": len(s),
                          "prob_media": s["prob_pick"].mean(), "acierto_real": s["acierto_pick"].mean()})
    r = pd.DataFrame(filas)
    r["diferencia"] = r["acierto_real"] - r["prob_media"]
    return r


def por_calificacion(pred: pd.DataFrame, por_metrica: bool = True) -> pd.DataFrame:
    claves = ["metrica", "calificacion"] if por_metrica else ["calificacion"]
    filas = []
    for clave, g in pred.groupby(claves, sort=False):
        clave = clave if isinstance(clave, tuple) else (clave,)
        filas.append({**dict(zip(claves, clave)), "patas": len(g), "prob_poisson_media": g["prob_poisson"].mean(),
                      "acierto_real": g["acierto"].mean(), "cuota_justa_modelo": 1 / g["prob_poisson"].mean(),
                      "cuota_minima_rentable": 1 / g["acierto"].mean() if g["acierto"].mean() > 0 else np.nan})
    r = pd.DataFrame(filas)
    r["diferencia"] = r["acierto_real"] - r["prob_poisson_media"]
    r["calificacion"] = pd.Categorical(r["calificacion"], ORDEN_CAL, ordered=True)
    if por_metrica:
        r["metrica"] = pd.Categorical(r["metrica"], METRICAS, ordered=True)
    return r.sort_values(claves).reset_index(drop=True)


def veredicto(pk: pd.DataFrame, pred: pd.DataFrame) -> pd.DataFrame:
    """Una fila por metrica: lo esencial para decidir si sus porcentajes son confiables."""
    filas = []
    for met in METRICAS:
        g, s = pk[pk["metrica"] == met], pred[pred["metrica"] == met]
        cal = calibracion(g)
        cal = cal[cal["picks"] >= 100]          # un tramo con pocos picks no dice nada del modelo
        exc = s[s["calificacion"] == "Excelente"]
        bue = s[s["calificacion"] == "Buena"]
        alto = g[g["prob_pick"] >= 0.8]
        res = resumen_mercados(g)
        filas.append({"metrica": met, "nombre": mo.NOMBRES[met],
                      "partidos": g.drop_duplicates(["liga", "fecha", "local", "visitante"]).shape[0],
                      "acierto_picks": g["acierto_pick"].mean(), "sin_modelo": ingenuo(g),
                      "habilidad_brier": np.average(res["habilidad_brier"], weights=res["partidos"]),
                      "desvio_max_calibracion": cal["diferencia"].abs().max(),
                      "desvio_medio_calibracion": np.average(cal["diferencia"], weights=cal["picks"]),
                      "dice_80+": alto["prob_pick"].mean(), "acierta_80+": alto["acierto_pick"].mean(), "picks_80+": len(alto),
                      "excelente_dice": exc["prob_poisson"].mean(), "excelente_acierta": exc["acierto"].mean(), "patas_excelente": len(exc),
                      "buena_dice": bue["prob_poisson"].mean(), "buena_acierta": bue["acierto"].mean(), "patas_buena": len(bue)})
    r = pd.DataFrame(filas)
    d = r["desvio_max_calibracion"]
    r["veredicto"] = np.where(d <= 0.03, "Confiable", np.where(d <= 0.06, "Aceptable (±3–6 pts)", "No confiable"))
    return r


def mayor_numero(pred: pd.DataFrame) -> pd.DataFrame:
    """Mercado 'quien hace mas': calibracion de las tres opciones."""
    g = pred[pred["grupo"] == "Mayor número"]
    filas = []
    for (met, mercado), s in g.groupby(["metrica", "mercado"], sort=False):
        for lo, hi in [(0, 0.3), (0.3, 0.4), (0.4, 0.5), (0.5, 0.6), (0.6, 0.7), (0.7, 1.01)]:
            t = s[(s["prob_poisson"] >= lo) & (s["prob_poisson"] < hi)]
            if len(t) >= 20:
                filas.append({"metrica": met, "mercado": mercado, "tramo_prob": f"{lo:.0%}–{min(hi, 1):.0%}", "patas": len(t),
                              "prob_media": t["prob_poisson"].mean(), "acierto_real": t["acierto"].mean()})
    r = pd.DataFrame(filas)
    r["diferencia"] = r["acierto_real"] - r["prob_media"]
    return r


# ------------------------------------------------------------------ excel
PCT = {"over_real_pct", "pct_picks_over", "acierto_modelo", "acierto_sin_modelo", "prob_media_pick", "mejora_vs_sin_modelo",
       "habilidad_brier", "prob_media", "acierto_real", "diferencia", "prob_poisson_media", "acierto_picks", "sin_modelo",
       "desvio_max_calibracion", "desvio_medio_calibracion", "dice_80+", "acierta_80+", "excelente_dice", "excelente_acierta",
       "buena_dice", "buena_acierta", "prob_poisson", "tasa_ult5"}


def escribir_excel(ruta: str, hojas: dict, notas: dict):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df = df.copy()
            if "fecha" in df.columns:
                df["fecha"] = pd.to_datetime(df["fecha"]).dt.strftime("%d/%m/%Y")
            df.to_excel(xw, sheet_name=nombre, index=False, startrow=2)
            ws = xw.sheets[nombre]
            ws["A1"] = notas.get(nombre, "")
            ws["A1"].font = Font(name="Arial", size=10, italic=True)
            for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
                for c in row:
                    c.font = Font(name="Arial", size=10, bold=(c.row == 3))
                    if c.row == 3:
                        c.fill = PatternFill("solid", fgColor="DDEBF7")
                        c.alignment = Alignment(wrap_text=True, vertical="center")
            for j, col in enumerate(df.columns, start=1):
                letra = get_column_letter(j)
                if col in PCT:
                    for c in ws[letra][3:]:
                        c.number_format = "0.0%"
                elif df[col].dtype.kind == "f":
                    for c in ws[letra][3:]:
                        c.number_format = "0.00"
                ancho = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) if len(df) else len(str(col))
                ws.column_dimensions[letra].width = min(max(10, ancho + 2), 30)
            ws.freeze_panes = "A4"


NOTAS = {
    "veredicto": "Una fila por métrica. desvio_max_calibracion = mayor diferencia entre lo que dice el modelo y lo que acierta en algún tramo de probabilidad (50–60, …, 90–100) con al menos 100 picks. Confiable ≤ 3 pts · Aceptable ≤ 6 · No confiable > 6. Solo partidos con ≥ 5 partidos en su condición para ambos equipos.",
    "calibracion": "Picks (lado con prob ≥ 50%) agrupados por la probabilidad que dio el modelo, por métrica. Calibrado = prob_media ≈ acierto_real. diferencia negativa = se pasa de confiado; positiva = se queda corto.",
    "calibracion_grupo": "Lo mismo, por métrica y familia de mercado (Total / Local / Visitante / Ambos anotan).",
    "calificacion": "Todas las patas (Over y Under) según la calificación de Kuota (0.6 × Poisson + 0.4 × últimos 5 de cada equipo), por métrica. cuota_minima_rentable = 1 / acierto real.",
    "calificacion_total": "Calificación de Kuota sumando todas las métricas.",
    "resumen": "Una fila por métrica y línea. c = línea central del partido (round(λ) − 0.5), c−1 = una línea abajo, etc. El modelo 'elige' el lado con prob ≥ 50%. acierto_sin_modelo = apostar siempre al lado más frecuente de esa línea en cada liga. habilidad_brier > 0 = mejora al promedio de liga.",
    "mayor_numero": "Mercado 'quién hace más' (tiros, corners, faltas, amarillas, tiros a puerta): probabilidad del modelo vs. acierto real por tramo.",
    "datos": "Efecto de la cantidad de datos: partidos donde algún equipo tenía menos de 5 partidos en su condición.",
}


def correr(rutas: list, metricas: list = METRICAS, procesos: int = 6) -> dict:
    with Pool(min(procesos, len(rutas))) as pool:
        partes = pool.map(predecir_liga, [(r, metricas) for r in rutas])
    pred = pd.concat(partes, ignore_index=True)
    ok = pred[pred["datos_suficientes"]]
    pk_ok, pk_all = picks(ok), picks(pred)

    def fila_datos(nombre, g):
        return {"grupo_datos": nombre, "partidos": g.drop_duplicates(["liga", "fecha", "local", "visitante"]).shape[0],
                "acierto_modelo": g["acierto_pick"].mean(), "acierto_sin_modelo": ingenuo(g),
                "prob_media_pick": g["prob_pick"].mean(), "desvio_max_calibracion": calibracion(g)["diferencia"].abs().max()}
    datos = pd.DataFrame([fila_datos("Ambos equipos con ≥ 5 partidos en su condición (usado en el resto de hojas)", pk_ok),
                          fila_datos("Algún equipo con < 5 partidos en su condición (pocos datos)", pk_all[~pk_all["datos_suficientes"]]),
                          fila_datos("Todos", pk_all)])
    return {"pred": pred, "ok": ok, "pk_ok": pk_ok, "pk_all": pk_all,
            "hojas": {"veredicto": veredicto(pk_ok, ok), "calibracion": calibracion(pk_ok, ["metrica"]),
                      "calibracion_grupo": calibracion(pk_ok, ["metrica", "grupo"]), "calificacion": por_calificacion(ok),
                      "calificacion_total": por_calificacion(ok, por_metrica=False), "resumen": resumen_mercados(pk_ok),
                      "mayor_numero": mayor_numero(ok), "datos": datos}}


if __name__ == "__main__":
    rutas = sys.argv[1:] or [f"datos/bbdd_{lg}.csv" for lg in LIGAS]
    r = correr(rutas)
    escribir_excel("backtest_kuota.xlsx", r["hojas"], NOTAS)
    r["pred"].to_csv("backtest_predicciones.csv", index=False)
    pd.set_option("display.width", 250)
    print(f"K_ENCOGE = {mo.K_ENCOGE} · partidos con predicción: {r['pred'].drop_duplicates(['liga', 'fecha', 'local', 'visitante']).shape[0]}")
    print(r["hojas"]["veredicto"].to_string(index=False))
    print("\nEscrito backtest_kuota.xlsx y backtest_predicciones.csv")